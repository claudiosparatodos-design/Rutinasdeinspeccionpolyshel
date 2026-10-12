#!/usr/bin/env python3
"""
Convierte los formatos originales de Excel (carpeta formatos/) en plantillas
que la página web puede dibujar tal cual (plantillas/<id>.json) y copia el
.xlsx original (plantillas/<id>.xlsx) para exportar llenando las mismas celdas.

Uso:
    python3 tools/convertir_plantillas.py

Requiere: openpyxl y Pillow.

El JSON contiene, por cada hoja: anchos de columna y altos de fila en
píxeles, celdas con su texto y estilo, celdas combinadas, imágenes y formas
dibujadas (cuadros de texto, casillas, flechas) y la configuración de página.
"""
import json
import math
import os
import re
import shutil
import sys
import zipfile
import colorsys
import posixpath
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string, range_boundaries
from openpyxl.cell.rich_text import CellRichText, TextBlock
from PIL import Image

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEN = os.path.join(RAIZ, 'formatos')
DESTINO = os.path.join(RAIZ, 'plantillas')

# id de plantilla -> archivo original
PLANTILLAS = {
    'consumo-agua': 'FORMATO CONSUMO DE AGUA.xlsx',
    'lecturas-cfe': 'FORMATO DE RECOLECCION DE LECTURAS.xlsx',
    'presiones': 'REGISTRO DE PRESIONES Y TEMPERATURA.xlsx',
    'servicios-turnos': 'RUTINA INSPECCION SERVICIOS PARA TURNOS.xlsx',
    'paro-arranque': 'RUTINAS PARO-ARRANQUE PLANTA.xlsx',
}

NS = {
    'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'm': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
}
EMU_PX = 9525.0
MDW = 7  # ancho máximo de dígito de Calibri 11 en píxeles

INDEXADOS = [
    '000000', 'FFFFFF', 'FF0000', '00FF00', '0000FF', 'FFFF00', 'FF00FF', '00FFFF',
    '000000', 'FFFFFF', 'FF0000', '00FF00', '0000FF', 'FFFF00', 'FF00FF', '00FFFF',
    '800000', '008000', '000080', '808000', '800080', '008080', 'C0C0C0', '808080',
    '9999FF', '993366', 'FFFFCC', 'CCFFFF', '660066', 'FF8080', '0066CC', 'CCCCFF',
    '000080', 'FF00FF', 'FFFF00', '00FFFF', '800080', '800000', '008080', '0000FF',
    '00CCFF', 'CCFFFF', 'CCFFCC', 'FFFF99', '99CCFF', 'FF99CC', 'CC99FF', 'FFCC99',
    '3366FF', '33CCCC', '99CC00', 'FFCC00', 'FF9900', 'FF6600', '666699', '969696',
    '003366', '339966', '003300', '333300', '993300', '993366', '333399', '333333',
]


# ----------------------------------------------------------------- colores
def leer_tema(z):
    try:
        raiz = ET.fromstring(z.read('xl/theme/theme1.xml'))
    except KeyError:
        return {}
    esquema = raiz.find('.//a:clrScheme', NS)
    colores = {}
    for hijo in esquema:
        nombre = hijo.tag.split('}')[1]
        c = hijo[0]
        if c.tag.endswith('srgbClr'):
            colores[nombre] = c.get('val')
        elif c.tag.endswith('sysClr'):
            colores[nombre] = c.get('lastClr', '000000')
    return colores


def tema_por_indice(tema, i):
    orden = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
    return tema.get(orden[i]) if 0 <= i < len(orden) else None


def aplicar_tinte(hexrgb, tinte):
    if not tinte:
        return hexrgb
    r, g, b = (int(hexrgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    if tinte < 0:
        l = l * (1 + tinte)
    else:
        l = l * (1 - tinte) + tinte
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))


def color_openpyxl(color, tema, defecto=None):
    if color is None:
        return defecto
    try:
        if color.type == 'rgb' and isinstance(color.rgb, str):
            v = color.rgb[-6:]
            if color.rgb.startswith('00') and len(color.rgb) == 8 and color.rgb == '00000000':
                return defecto
            return '#' + aplicar_tinte(v, color.tint)
        if color.type == 'theme':
            base = tema_por_indice(tema, color.theme)
            if base:
                return '#' + aplicar_tinte(base, color.tint)
        if color.type == 'indexed':
            if color.indexed in (64, 65):
                return defecto
            if color.indexed < len(INDEXADOS):
                return '#' + aplicar_tinte(INDEXADOS[color.indexed], color.tint)
    except Exception:
        pass
    return defecto


def color_drawingml(nodo, tema):
    """Lee srgbClr/schemeClr/sysClr con modificadores (lumMod, lumOff, shade, tint)."""
    if nodo is None:
        return None
    c = None
    for hijo in nodo:
        etiqueta = hijo.tag.split('}')[1]
        if etiqueta in ('srgbClr', 'schemeClr', 'sysClr', 'prstClr', 'scrgbClr'):
            c = hijo
            break
    if c is None:
        return None
    etiqueta = c.tag.split('}')[1]
    if etiqueta == 'srgbClr':
        base = c.get('val')
    elif etiqueta == 'sysClr':
        base = c.get('lastClr', '000000')
    elif etiqueta == 'schemeClr':
        val = c.get('val')
        alias = {'tx1': 'dk1', 'bg1': 'lt1', 'tx2': 'dk2', 'bg2': 'lt2'}
        base = tema.get(alias.get(val, val), '000000')
    elif etiqueta == 'prstClr':
        base = {'black': '000000', 'white': 'FFFFFF', 'red': 'FF0000'}.get(c.get('val'), '000000')
    else:
        base = '000000'
    r, g, b = (int(base[i:i + 2], 16) / 255 for i in (0, 2, 4))
    for mod in c:
        m = mod.tag.split('}')[1]
        v = int(mod.get('val', '100000')) / 100000
        if m in ('lumMod', 'lumOff'):
            h, l, s = colorsys.rgb_to_hls(r, g, b)
            l = l * v if m == 'lumMod' else min(1, l + v)
            r, g, b = colorsys.hls_to_rgb(h, l, s)
        elif m == 'shade':
            r, g, b = r * v, g * v, b * v
        elif m == 'tint':
            r, g, b = r + (1 - r) * (1 - v), g + (1 - g) * (1 - v), b + (1 - b) * (1 - v)
    return '#%02X%02X%02X' % (round(min(1, r) * 255), round(min(1, g) * 255), round(min(1, b) * 255))


# ----------------------------------------------------------------- estilos
BORDES = {
    'thin': (1, 'solid'), 'medium': (2, 'solid'), 'thick': (3, 'solid'), 'double': (3, 'double'),
    'hair': (1, 'solid'), 'dotted': (1, 'dotted'), 'dashed': (1, 'dashed'), 'mediumDashed': (2, 'dashed'),
    'dashDot': (1, 'dashed'), 'mediumDashDot': (2, 'dashed'), 'dashDotDot': (1, 'dotted'),
    'mediumDashDotDot': (2, 'dotted'), 'slantDashDot': (2, 'dashed'),
}


def estilo_celda(c, tema, fuente_tema):
    f = c.font
    e = {}
    nombre = f.name or fuente_tema
    if f.scheme == 'minor':
        nombre = fuente_tema
    e['ff'] = nombre
    e['fs'] = float(f.sz or 11)
    if f.b:
        e['b'] = 1
    if f.i:
        e['i'] = 1
    if f.u:
        e['u'] = 1
    if f.strike:
        e['st'] = 1
    col = color_openpyxl(f.color, tema, None)
    if col and col.upper() != '#000000':
        e['c'] = col
    fill = c.fill
    if fill is not None and fill.fill_type == 'solid':
        bg = color_openpyxl(fill.fgColor, tema, None)
        if bg is None:
            bg = color_openpyxl(fill.bgColor, tema, None)
        if bg and bg.upper() != '#FFFFFF':
            e['bg'] = bg
        elif bg:
            e['bg'] = '#FFFFFF'
    for lado, clave in (('left', 'bl'), ('right', 'br'), ('top', 'bt'), ('bottom', 'bb')):
        b = getattr(c.border, lado)
        if b is not None and b.style:
            ancho, tipo = BORDES.get(b.style, (1, 'solid'))
            colb = color_openpyxl(b.color, tema, '#000000') or '#000000'
            e[clave] = [ancho, tipo, colb]
    a = c.alignment
    if a.horizontal and a.horizontal != 'general':
        e['ha'] = a.horizontal
    if a.vertical:
        e['va'] = a.vertical
    if a.wrap_text:
        e['w'] = 1
    if a.indent:
        e['in'] = int(a.indent)
    if a.text_rotation:
        e['rot'] = int(a.text_rotation)
    if a.shrink_to_fit:
        e['sh'] = 1
    if c.number_format and c.number_format != 'General':
        e['nf'] = c.number_format
    return e


def texto_valor(v):
    if v is None:
        return None
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return ('%.6f' % v).rstrip('0').rstrip('.')
    if hasattr(v, 'strftime'):
        return v.strftime('%d/%m/%Y')
    return str(v)


def runs_rich(v, tema, fuente_tema):
    runs = []
    for parte in v:
        if isinstance(parte, TextBlock):
            f = parte.font
            r = {'t': parte.text}
            if f is not None:
                if f.b:
                    r['b'] = 1
                if f.i:
                    r['i'] = 1
                if f.u:
                    r['u'] = 1
                if f.sz:
                    r['fs'] = float(f.sz)
                if f.rFont:
                    r['ff'] = f.rFont
                colr = color_openpyxl(f.color, tema, None)
                if colr and colr.upper() != '#000000':
                    r['c'] = colr
            runs.append(r)
        else:
            runs.append({'t': str(parte)})
    return runs


# ----------------------------------------------------------------- dibujos
def leer_rels(z, ruta):
    carpeta, archivo = posixpath.split(ruta)
    rels = posixpath.join(carpeta, '_rels', archivo + '.rels')
    try:
        raiz = ET.fromstring(z.read(rels))
    except KeyError:
        return {}
    out = {}
    for r in raiz:
        destino = r.get('Target')
        if r.get('TargetMode') == 'External':
            continue
        out[r.get('Id')] = (r.get('Type').split('/')[-1], posixpath.normpath(posixpath.join(carpeta, destino)))
    return out


def rutas_hojas(z):
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    rels = leer_rels(z, 'xl/workbook.xml')
    out = []
    for s in wb.find('m:sheets', NS):
        rid = s.get('{%s}id' % NS['r'])
        out.append((s.get('name'), rels[rid][1]))
    return out


def ruta_path(path_el, w_esc, h_esc):
    d = []
    for cmd in path_el:
        t = cmd.tag.split('}')[1]
        pts = [(float(p.get('x')) * w_esc, float(p.get('y')) * h_esc) for p in cmd.findall('a:pt', NS)]
        if t == 'moveTo':
            d.append('M%.2f %.2f' % pts[0])
        elif t == 'lnTo':
            d.append('L%.2f %.2f' % pts[0])
        elif t == 'cubicBezTo':
            d.append('C' + ' '.join('%.2f %.2f' % p for p in pts))
        elif t == 'quadBezTo':
            d.append('Q' + ' '.join('%.2f %.2f' % p for p in pts))
        elif t == 'close':
            d.append('Z')
    return ' '.join(d)


def leer_dibujo(z, ruta_dib, tema, x_col, y_fila, id_plantilla, n_hoja, fuente_tema):
    raiz = ET.fromstring(z.read(ruta_dib))
    rels = leer_rels(z, ruta_dib)
    items = []

    def posicion(ancla):
        etiqueta = ancla.tag.split('}')[1]
        if etiqueta == 'absoluteAnchor':
            pos = ancla.find('xdr:pos', NS)
            ext = ancla.find('xdr:ext', NS)
            x, y = int(pos.get('x')) / EMU_PX, int(pos.get('y')) / EMU_PX
            return x, y, int(ext.get('cx')) / EMU_PX, int(ext.get('cy')) / EMU_PX, None
        fr = ancla.find('xdr:from', NS)
        c0, co0 = int(fr.find('xdr:col', NS).text), int(fr.find('xdr:colOff', NS).text)
        r0, ro0 = int(fr.find('xdr:row', NS).text), int(fr.find('xdr:rowOff', NS).text)
        x0, y0 = x_col(c0) + co0 / EMU_PX, y_fila(r0) + ro0 / EMU_PX
        if etiqueta == 'oneCellAnchor':
            ext = ancla.find('xdr:ext', NS)
            return x0, y0, int(ext.get('cx')) / EMU_PX, int(ext.get('cy')) / EMU_PX, (c0, r0)
        to = ancla.find('xdr:to', NS)
        c1, co1 = int(to.find('xdr:col', NS).text), int(to.find('xdr:colOff', NS).text)
        r1, ro1 = int(to.find('xdr:row', NS).text), int(to.find('xdr:rowOff', NS).text)
        x1, y1 = x_col(c1) + co1 / EMU_PX, y_fila(r1) + ro1 / EMU_PX
        return x0, y0, max(1, x1 - x0), max(1, y1 - y0), (c0, r0)

    n_img = 0
    for ancla in raiz:
        x, y, w, h, desde = posicion(ancla)
        pic = ancla.find('xdr:pic', NS)
        sp = ancla.find('xdr:sp', NS)
        base = {'x': round(x, 1), 'y': round(y, 1), 'w': round(w, 1), 'h': round(h, 1)}
        if desde:
            base['desde'] = [desde[0], desde[1]]
        if pic is not None:
            blip = pic.find('.//a:blip', NS)
            rid = blip.get('{%s}embed' % NS['r']) if blip is not None else None
            if not rid or rid not in rels:
                continue
            datos = z.read(rels[rid][1])
            n_img += 1
            from io import BytesIO
            im = Image.open(BytesIO(datos))
            im = im.convert('RGBA')
            # Recorte aplicado en Excel (a:srcRect, en milésimas de porcentaje)
            src = pic.find('.//a:srcRect', NS)
            if src is not None:
                W, H = im.size
                l, t = int(src.get('l', 0)) / 100000, int(src.get('t', 0)) / 100000
                r, b = int(src.get('r', 0)) / 100000, int(src.get('b', 0)) / 100000
                caja = (round(W * l), round(H * t), round(W * (1 - r)), round(H * (1 - b)))
                if caja[2] > caja[0] and caja[3] > caja[1]:
                    im = im.crop(caja)
            # Reducir a 3x el tamaño en pantalla para que pese poco y se vea nítida
            lado = max(1, int(max(w, h) * 3))
            if max(im.size) > lado:
                esc = lado / max(im.size)
                im = im.resize((max(1, int(im.size[0] * esc)), max(1, int(im.size[1] * esc))), Image.LANCZOS)
            import hashlib
            buf = BytesIO()
            im.save(buf, format='PNG', optimize=True)
            huella = hashlib.md5(buf.getvalue()).hexdigest()[:10]
            nombre = 'img-%s.png' % huella
            destino = os.path.join(DESTINO, 'img', nombre)
            if not os.path.exists(destino):
                with open(destino, 'wb') as fimg:
                    fimg.write(buf.getvalue())
            item = dict(base, tipo='img', src='plantillas/img/' + nombre)
            items.append(item)
        elif sp is not None:
            nv = sp.find('xdr:nvSpPr/xdr:cNvPr', NS)
            sppr = sp.find('xdr:spPr', NS)
            estilo = sp.find('xdr:style', NS)
            item = dict(base, tipo='forma', id=nv.get('id'), nombre=nv.get('name'))
            prst = sppr.find('a:prstGeom', NS)
            cust = sppr.find('a:custGeom', NS)
            if cust is not None:
                rutas = []
                for p in cust.findall('a:pathLst/a:path', NS):
                    pw, ph = float(p.get('w', 1) or 1), float(p.get('h', 1) or 1)
                    rutas.append(ruta_path(p, w / pw, h / ph))
                item['geom'] = 'path'
                item['d'] = ' '.join(rutas)
            else:
                item['geom'] = prst.get('prst') if prst is not None else 'rect'
            # Relleno
            relleno = None
            if sppr.find('a:noFill', NS) is not None:
                relleno = None
            elif sppr.find('a:solidFill', NS) is not None:
                relleno = color_drawingml(sppr.find('a:solidFill', NS), tema)
            elif estilo is not None:
                fr_ = estilo.find('a:fillRef', NS)
                if fr_ is not None and fr_.get('idx') != '0':
                    relleno = color_drawingml(fr_, tema)
            if relleno:
                item['relleno'] = relleno
            # Línea
            ln = sppr.find('a:ln', NS)
            linea, grosor = None, 1.0
            if ln is not None and ln.find('a:noFill', NS) is not None:
                linea = None
            else:
                if ln is not None and ln.find('a:solidFill', NS) is not None:
                    linea = color_drawingml(ln.find('a:solidFill', NS), tema)
                elif estilo is not None:
                    lr = estilo.find('a:lnRef', NS)
                    if lr is not None and lr.get('idx') != '0':
                        linea = color_drawingml(lr, tema)
                if ln is not None and ln.get('w'):
                    grosor = int(ln.get('w')) / 12700 * 1.333
                elif linea:
                    grosor = 1.0
            if linea:
                item['linea'] = linea
                item['grosor'] = round(grosor, 2)
            # Texto
            color_fuente = '#000000'
            if estilo is not None:
                fref = estilo.find('a:fontRef', NS)
                if fref is not None:
                    color_fuente = color_drawingml(fref, tema) or color_fuente
            tx = sp.find('xdr:txBody', NS)
            parrafos = []
            if tx is not None:
                body = tx.find('a:bodyPr', NS)
                if body is not None and body.get('anchor'):
                    item['ancla'] = body.get('anchor')
                for p in tx.findall('a:p', NS):
                    ppr = p.find('a:pPr', NS)
                    par = {'runs': []}
                    if ppr is not None and ppr.get('algn'):
                        par['algn'] = ppr.get('algn')
                    for r in p.findall('a:r', NS):
                        rpr = r.find('a:rPr', NS)
                        t = r.find('a:t', NS)
                        run = {'t': t.text or ''}
                        if rpr is not None:
                            if rpr.get('sz'):
                                run['fs'] = int(rpr.get('sz')) / 100
                            if rpr.get('b') == '1':
                                run['b'] = 1
                            if rpr.get('i') == '1':
                                run['i'] = 1
                            cr = color_drawingml(rpr.find('a:solidFill', NS), tema) if rpr.find('a:solidFill', NS) is not None else None
                            if cr:
                                run['c'] = cr
                        par['runs'].append(run)
                    end = p.find('a:endParaRPr', NS)
                    if end is not None and end.get('sz'):
                        par['fs'] = int(end.get('sz')) / 100
                    parrafos.append(par)
            if any(r['t'].strip() for p in parrafos for r in p['runs']):
                item['texto'] = parrafos
                item['colorTexto'] = color_fuente
            items.append(item)
    return items


# ----------------------------------------------------------------- hoja
def convertir_hoja(z, ws, ruta_hoja, tema, fuente_tema, id_plantilla, n_hoja, estilos, indice_estilos):
    # Anchos de columna
    base = ws.sheet_format.baseColWidth or 8
    defecto_ancho = ws.sheet_format.defaultColWidth
    px_defecto = int(defecto_ancho * MDW + 0.5) if defecto_ancho else int(base * MDW + 10)
    anchos = {}
    ocultas = set()
    for clave, dim in ws.column_dimensions.items():
        mn, mx = dim.min or column_index_from_string(clave), dim.max or column_index_from_string(clave)
        for c in range(mn, mx + 1):
            if dim.hidden:
                ocultas.add(c)
            if dim.width:
                anchos[c] = int(dim.width * MDW + 0.5)
    alto_defecto = ws.sheet_format.defaultRowHeight or 15
    altos = {}
    filas_ocultas = set()
    for r, dim in ws.row_dimensions.items():
        if dim.hidden:
            filas_ocultas.add(r)
        if dim.height is not None:
            altos[r] = dim.height

    def px_col(c):
        return 0 if c in ocultas else anchos.get(c, px_defecto)

    def px_fila(r):
        return 0 if r in filas_ocultas else round(altos.get(r, alto_defecto) * 4 / 3, 2)

    def x_col(c0):  # índice base 0
        return sum(px_col(c) for c in range(1, c0 + 1))

    def y_fila(r0):
        return sum(px_fila(r) for r in range(1, r0 + 1))

    # Celdas con contenido o estilo visible
    celdas = {}
    max_c, max_r = 1, 1
    for fila in ws.iter_rows():
        for c in fila:
            if not hasattr(c, 'column') or c.column is None:
                continue
            e = estilo_celda(c, tema, fuente_tema)
            visible = c.value is not None or any(k in e for k in ('bl', 'br', 'bt', 'bb', 'bg'))
            clave_e = json.dumps(e, sort_keys=True)
            if clave_e not in indice_estilos:
                indice_estilos[clave_e] = len(estilos)
                estilos.append(e)
            s = indice_estilos[clave_e]
            dato = {'s': s}
            if isinstance(c.value, CellRichText):
                dato['r'] = runs_rich(c.value, tema, fuente_tema)
            elif c.value is not None:
                dato['v'] = texto_valor(c.value)
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    dato['n'] = 1
            if c.value is not None or s != 0 or visible:
                celdas[c.coordinate] = dato
            if visible:
                max_c, max_r = max(max_c, c.column), max(max_r, c.row)

    combinadas = [str(m) for m in ws.merged_cells.ranges]
    for m in ws.merged_cells.ranges:
        max_c, max_r = max(max_c, m.max_col), max(max_r, m.max_row)

    dibujos = []
    rels = leer_rels(z, ruta_hoja)
    for tipo, destino in rels.values():
        if tipo == 'drawing':
            dibujos += leer_dibujo(z, destino, tema, x_col, y_fila, id_plantilla, n_hoja, fuente_tema)

    # Extender el área hasta donde lleguen los dibujos
    for d in dibujos:
        while x_col(max_c) < d['x'] + d['w'] - 2:
            max_c += 1
        while y_fila(max_r) < d['y'] + d['h'] - 2:
            max_r += 1

    cols = [px_col(c) for c in range(1, max_c + 1)]
    filas = [px_fila(r) for r in range(1, max_r + 1)]
    # Descartar celdas fuera del área
    celdas = {k: v for k, v in celdas.items()
              if column_index_from_string(re.match(r'[A-Z]+', k).group()) <= max_c and int(re.search(r'\d+', k).group()) <= max_r}

    ps = ws.page_setup
    pm = ws.page_margins
    ajustar = bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage)
    pagina = {
        'orientacion': ps.orientation or 'portrait',
        'escala': int(ps.scale) if ps.scale else 100,
        'ajustar': ajustar,
        'anchoPaginas': int(ps.fitToWidth) if ps.fitToWidth is not None else 1,
        'altoPaginas': int(ps.fitToHeight) if ps.fitToHeight is not None else 1,
        'margenes': {'l': pm.left, 'r': pm.right, 't': pm.top, 'b': pm.bottom},
        'centrarH': bool(ws.print_options.horizontalCentered),
        'centrarV': bool(ws.print_options.verticalCentered),
    }
    return {
        'nombre': ws.title,
        'cols': cols,
        'filas': filas,
        'ancho': round(sum(cols), 1),
        'alto': round(sum(filas), 1),
        'celdas': celdas,
        'combinadas': combinadas,
        'dibujos': dibujos,
        'pagina': pagina,
    }


def convertir(id_plantilla, archivo):
    ruta = os.path.join(ORIGEN, archivo)
    z = zipfile.ZipFile(ruta)
    tema = leer_tema(z)
    try:
        t = ET.fromstring(z.read('xl/theme/theme1.xml'))
        fuente_tema = t.find('.//a:minorFont/a:latin', NS).get('typeface') or 'Calibri'
    except Exception:
        fuente_tema = 'Calibri'
    wb = openpyxl.load_workbook(ruta, rich_text=True)
    rutas = dict(rutas_hojas(z))
    estilos, indice = [], {}
    # Estilo 0 = celda sin formato
    indice[json.dumps({'ff': fuente_tema, 'fs': 11.0}, sort_keys=True)] = 0
    estilos.append({'ff': fuente_tema, 'fs': 11.0})
    hojas = []
    for n, ws in enumerate(wb.worksheets):
        hojas.append(convertir_hoja(z, ws, rutas[ws.title], tema, fuente_tema, id_plantilla, n + 1, estilos, indice))
    salida = {'id': id_plantilla, 'archivo': archivo, 'fuente': fuente_tema, 'estilos': estilos, 'hojas': hojas}
    with open(os.path.join(DESTINO, id_plantilla + '.json'), 'w', encoding='utf8') as f:
        json.dump(salida, f, ensure_ascii=False, separators=(',', ':'))
    shutil.copyfile(ruta, os.path.join(DESTINO, id_plantilla + '.xlsx'))
    print('%-18s %d hojas, %d estilos' % (id_plantilla, len(hojas), len(estilos)))


if __name__ == '__main__':
    os.makedirs(os.path.join(DESTINO, 'img'), exist_ok=True)
    for f in os.listdir(os.path.join(DESTINO, 'img')):
        os.remove(os.path.join(DESTINO, 'img', f))
    for k, v in PLANTILLAS.items():
        convertir(k, v)
